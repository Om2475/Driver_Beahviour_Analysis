import shutil
from docx import Document
from pathlib import Path


def delete_paragraph(paragraph):
    p = paragraph._element
    p.getparent().remove(p)
    paragraph._p = paragraph._element = None


def replace_section(doc: Document, start_title: str, end_title: str, new_text: str) -> bool:
    paragraphs = list(doc.paragraphs)
    start_idx = None
    end_idx = None
    for i, p in enumerate(paragraphs):
        if p.text and p.text.strip().lower().startswith(start_title.lower()):
            start_idx = i
            break
    for i, p in enumerate(paragraphs):
        if end_title and p.text and p.text.strip().lower().startswith(end_title.lower()):
            end_idx = i
            break

    if start_idx is None:
        return False

    # If end title not found, we'll remove until end
    if end_idx is None:
        to_remove = paragraphs[start_idx + 1 :]
    else:
        to_remove = paragraphs[start_idx + 1 : end_idx]

    for p in to_remove:
        try:
            delete_paragraph(p)
        except Exception:
            pass

    # Insert new text after start_idx paragraph
    insert_after = doc.paragraphs[start_idx]._element
    for line in new_text.strip().split("\n"):
        new_p = Document().add_paragraph(line)
        insert_after.addnext(new_p._p)
        insert_after = new_p._p

    return True


def insert_section_at_end(doc: Document, title: str, text: str):
    doc.add_paragraph(title)
    for line in text.strip().split("\n"):
        doc.add_paragraph(line)


def main():
    base = Path(__file__).resolve().parents[1]
    report = base / "DriverBehaviour_Report_Final.docx"
    if not report.exists():
        print(f"Report not found: {report}")
        return

    backup = report.with_name(report.stem + "_backup.docx")
    shutil.copy(report, backup)
    print(f"Backup created: {backup.name}")

    doc = Document(str(report))

    proposed_methodology = (
        "Proposed Methodology\n"
        "Objective: Identify recurring driving contexts from cleaned OBD-II telemetry and produce concise, interpretable trip summaries.\n\n"
        "Data: Use the cleaned, windowed feature table produced in feature engineering.\n\n"
        "Steps:\n"
        "1. Standardize and validate window records (alignment, sensor bounds, no missing values).\n"
        "2. Produce a compact representation for each window for grouping (e.g., PCA/UMAP or an autoencoder).\n"
        "3. Group windows into context states using an unsupervised method tuned for interpretability and within-context homogeneity.\n"
        "4. Apply short rolling-mode smoothing per trip to reduce spurious label flips but preserve true transitions.\n"
        "5. Compute a per-window confidence score combining group-fit, reconstruction error, and local density; flag low-confidence windows as anomalous.\n"
        "6. Profile each context via z-scored feature means and report the top distinguishing features.\n\n"
        "Outputs:\n"
        "- Window-level context labels (raw and smoothed), confidence scores, and anomaly flags.\n"
        "- Per-context profiles listing key elevated or reduced features.\n"
        "- Per-trip summaries: context proportions, mean confidence, anomaly ratio, mean dwell times, and transition counts.\n\n"
        "Validation:\n"
        "- Use holdout sampling and small perturbations to report stability (adjusted rand index / repeatability).\n"
        "- Provide clear diagnostics: within-context variance, transition stability, and anomaly ratio."
    )

    conclusion = (
        "Conclusion\n"
        "This pipeline converts cleaned OBD-II telemetry into clear, repeatable driving contexts and compact trip summaries for practical use.\n\n"
        "We prioritize data hygiene, simple unsupervised grouping, and conservative smoothing to avoid over-interpreting transient noise. The outputs — labeled windows, confidence scores, anomaly flags, and per-trip summaries — are designed for direct consumption in dashboards or manual review workflows.\n\n"
        "Limitations: Contexts are descriptive and should be validated against labeled events or driver feedback for operational deployment. Re-evaluate context definitions when sensor setups or vehicle types change.\n\n"
        "Suggested next steps: add a short appendix with representative context profiles and example time-series plots, and set up a lightweight human-review loop to confirm critical anomaly cases."
    )

    replaced_pm = replace_section(doc, "Proposed Methodology", "Conclusion", proposed_methodology)
    if not replaced_pm:
        insert_section_at_end(doc, "Proposed Methodology", proposed_methodology)

    replaced_conc = replace_section(doc, "Conclusion", "", conclusion)
    if not replaced_conc:
        insert_section_at_end(doc, "Conclusion", conclusion)

    updated = report.with_name(report.stem + "_updated.docx")
    doc.save(str(updated))
    print(f"Updated report saved: {updated.name}")


if __name__ == "__main__":
    main()
